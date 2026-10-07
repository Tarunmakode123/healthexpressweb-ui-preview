import openpyxl
import json
from pathlib import Path

wb = openpyxl.load_workbook(r'C:\Users\tarun\Downloads\HealthPackages.xlsx', data_only=True)

# 1. Inspect SexualWellness Sheet in detail
sw = wb['SexualWellness']
rows_sw = list(sw.iter_rows(values_only=True))

print("=== SEXUAL WELLNESS SHEET EXPLICIT PARSING ===")
header_group = rows_sw[0]  # Row 1 (0-indexed)
header_name = rows_sw[1]   # Row 2
mrp_row = rows_sw[38]      # Row 39 (MRP)

sw_packages = []
for col_idx in range(2, 9): # Columns C to I (col index 2 to 8)
    group_title = ''
    for prev in range(col_idx, -1, -1):
        if header_group[prev]:
            group_title = str(header_group[prev]).strip()
            break
            
    pkg_name = str(header_name[col_idx]).strip() if header_name[col_idx] else f"Package {col_idx}"
    mrp_val = mrp_row[col_idx] if col_idx < len(mrp_row) else None
    mrp = float(mrp_val) if mrp_val and str(mrp_val).replace('.', '', 1).isdigit() else 0
    
    # Collect parameters where cell == 'yes'
    params = []
    for r_idx in range(2, 38):
        r = rows_sw[r_idx]
        test_name = r[1]
        val = str(r[col_idx]).strip().lower() if col_idx < len(r) and r[col_idx] is not None else ''
        if test_name and val == 'yes':
            params.append(str(test_name).strip())
            
    # Format clean package name
    clean_group = group_title.replace('Packages', '').replace('(Premium)', '').strip()
    full_name = f"{clean_group} - {pkg_name}"
    if pkg_name in clean_group:
        full_name = clean_group
        
    sw_packages.append({
        'col_idx': col_idx,
        'group': group_title,
        'name': pkg_name,
        'full_name': full_name,
        'mrp': mrp,
        'params_count': len(params),
        'params': params
    })

for p in sw_packages:
    print(f"\nPackage Col {p['col_idx']}: [{p['group']}] -> [{p['name']}]")
    print(f"  Full Name: {p['full_name']}")
    print(f"  MRP: Rs. {p['mrp']}")
    print(f"  Included Parameters ({p['params_count']}): {p['params']}")

# 2. Inspect Preventive Care Sheet in detail
print("\n=== PREVENTIVE CARE SHEET EXPLICIT PARSING ===")
sp = wb['Preventive Care']
rows_sp = list(sp.iter_rows(values_only=True))
header_p = rows_sp[2]      # Row 3 (0-indexed)
mrp_row_p = rows_sp[43]    # Row 44 (MRP)

sp_packages = []
for col_idx in range(3, 12): # Columns D to L
    pkg_name = str(header_p[col_idx]).strip() if header_p[col_idx] else None
    if not pkg_name: continue
    
    mrp_val = mrp_row_p[col_idx] if col_idx < len(mrp_row_p) else None
    mrp = float(mrp_val) if mrp_val and str(mrp_val).replace('.', '', 1).isdigit() else 0
    
    params = []
    for r_idx in range(3, 43):
        r = rows_sp[r_idx]
        test_name = r[1]
        val = str(r[col_idx]).strip().lower() if col_idx < len(r) and r[col_idx] is not None else ''
        if test_name and val == 'yes':
            params.append(str(test_name).strip())
            
    sp_packages.append({
        'col_idx': col_idx,
        'name': pkg_name,
        'mrp': mrp,
        'params_count': len(params),
        'params': params
    })

for p in sp_packages:
    print(f"\nPreventive Package Col {p['col_idx']}: [{p['name']}]")
    print(f"  MRP: Rs. {p['mrp']}")
    print(f"  Included Parameters ({p['params_count']}): {p['params']}")
