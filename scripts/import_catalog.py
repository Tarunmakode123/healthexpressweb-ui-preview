import openpyxl
import re
import json
from pathlib import Path

def slugify(text):
    text = str(text).lower().strip()
    text = re.sub(r'[^\w\s-]', '', text)
    text = re.sub(r'[\s_-]+', '-', text)
    return text.strip('-')

slug_set = set()
services = []

stats = {
    'labspring': 0,
    'imaging': 0,
    'genomics': 0,
    'packages': 0,
    'missing_prices': 0,
    'missing_tat': 0,
    'missing_prep': 0,
    'missing_sample': 0,
    'duplicates_handled': 0
}

downloads_dir = Path(r'C:\Users\tarun\Downloads')

# 1. IMPORT LABSPRING
f1 = downloads_dir / 'LabSpring_Uploaded.xlsx'
if f1.exists():
    wb1 = openpyxl.load_workbook(f1, data_only=True)
    s1 = wb1['Table 1']
    rows1 = list(s1.iter_rows(values_only=True))
    for r in rows1[1:]:
        provider = str(r[0]).strip() if r[0] else 'LabSpring'
        cat_raw = str(r[1]).strip() if r[1] else 'Labs'
        desc = str(r[2]).strip() if r[2] and str(r[2]).lower() != 'none' else None
        code = str(r[3]).strip() if r[3] else None
        name = str(r[4]).strip() if r[4] else None
        if not name:
            continue

        mrp = float(r[5]) if r[5] is not None and str(r[5]).replace('.', '', 1).isdigit() else None
        b2b = float(r[6]) if r[6] is not None and str(r[6]).replace('.', '', 1).isdigit() else None
        disc_price = float(r[7]) if r[7] is not None and str(r[7]).replace('.', '', 1).isdigit() else mrp
        selling_price = disc_price if disc_price is not None else mrp

        if selling_price is None:
            stats['missing_prices'] += 1

        tat = str(r[8]).strip() if r[8] and str(r[8]).lower() != 'none' else None
        if not tat:
            stats['missing_tat'] += 1

        prep = str(r[9]).strip() if r[9] and str(r[9]).lower() != 'none' else None
        if not prep:
            stats['missing_prep'] += 1

        sample = str(r[10]).strip() if r[10] and str(r[10]).lower() != 'none' else None
        if not sample:
            stats['missing_sample'] += 1

        hc_raw = str(r[11]).strip().lower() if r[11] else ''
        home_collection = hc_raw in ['yes', 'y', 'true', 'available']

        code_str = code if code else ''
        base_slug = slugify(f"{name}-{provider}-{code_str}")
        slug = base_slug
        idx = 1
        while slug in slug_set:
            stats['duplicates_handled'] += 1
            slug = f"{base_slug}-{idx}"
            idx += 1
        slug_set.add(slug)

        disc_pct = None
        if mrp and selling_price and mrp > selling_price:
            pct = round(((mrp - selling_price) / mrp) * 100)
            if pct > 0:
                disc_pct = f"{pct}%"

        services.append({
            'id': slug,
            'slug': slug,
            'service_code': code,
            'name': name,
            'category_id': 'lab-tests',
            'category_name': 'Lab Tests',
            'subcategory': cat_raw,
            'provider': provider,
            'description': desc,
            'overview': desc,
            'mrp': mrp,
            'price': mrp,
            'selling_price': selling_price,
            'discount_price': selling_price,
            'b2b_price': b2b,
            'discount_percentage': disc_pct,
            'turnaround_time': tat or 'Contact for TAT',
            'patient_preparation': prep,
            'preparation': prep,
            'specimen_type': sample,
            'sample_type': sample or 'Standard Specimen',
            'home_collection_available': home_collection,
            'centre_visit_required': False,
            'service_type': 'lab',
            'source': 'LabSpring'
        })
        stats['labspring'] += 1

# 2. IMPORT IMAGING / RADIOLOGY
f2 = downloads_dir / 'Imaging_Radiology_MRP.xlsx'
if f2.exists():
    wb2 = openpyxl.load_workbook(f2, data_only=True)
    s2 = wb2['Imaging']
    rows2 = list(s2.iter_rows(values_only=True))
    for r in rows2[1:]:
        provider = str(r[0]).strip() if r[0] else 'Cadabams'
        cat_raw = str(r[1]).strip() if r[1] else 'Radiology'
        code = str(r[2]).strip() if r[2] else None
        name = str(r[3]).strip() if r[3] else None
        if not name:
            continue

        mrp = float(r[4]) if r[4] is not None and str(r[4]).replace('.', '', 1).isdigit() else None
        if mrp is None:
            stats['missing_prices'] += 1

        code_str = code if code else ''
        base_slug = slugify(f"{name}-{provider}-{code_str}")
        slug = base_slug
        idx = 1
        while slug in slug_set:
            stats['duplicates_handled'] += 1
            slug = f"{base_slug}-{idx}"
            idx += 1
        slug_set.add(slug)

        services.append({
            'id': slug,
            'slug': slug,
            'service_code': code,
            'name': name,
            'category_id': 'imaging',
            'category_name': 'Imaging & Radiology',
            'subcategory': cat_raw,
            'provider': provider,
            'description': f"Diagnostic {cat_raw} procedure performed by {provider}.",
            'overview': f"Diagnostic {cat_raw} procedure performed by {provider}.",
            'mrp': mrp,
            'price': mrp,
            'selling_price': mrp,
            'discount_price': mrp,
            'b2b_price': None,
            'discount_percentage': None,
            'turnaround_time': 'Same Day',
            'patient_preparation': 'Follow radiology procedure guidelines',
            'preparation': 'Follow radiology procedure guidelines',
            'specimen_type': 'Radiology Scan',
            'sample_type': 'Procedure / Scan',
            'home_collection_available': False,
            'centre_visit_required': True,
            'service_type': 'imaging',
            'source': 'Cadabams'
        })
        stats['imaging'] += 1

# 3. IMPORT GENOMICS
f3 = downloads_dir / 'Genomics_Test_Pricing.xlsx'
if f3.exists():
    wb3 = openpyxl.load_workbook(f3, data_only=True)
    s3 = wb3['Genomics Pricing']
    rows3 = list(s3.iter_rows(values_only=True))
    for r in rows3[1:]:
        cat_raw = str(r[0]).strip() if r[0] else 'Genetic Wellness'
        name = str(r[1]).strip() if r[1] else None
        if not name:
            continue

        mrp = float(r[2]) if r[2] is not None and str(r[2]).replace('.', '', 1).isdigit() else None
        tat = str(r[3]).strip() if r[3] and str(r[3]).lower() != 'none' else None
        sample = str(r[4]).strip() if r[4] and str(r[4]).lower() != 'none' else None
        prep = str(r[5]).strip() if r[5] and str(r[5]).lower() != 'none' else None
        desc = str(r[6]).strip() if r[6] and str(r[6]).lower() != 'none' else None

        base_slug = slugify(f"{name}-genomics")
        slug = base_slug
        idx = 1
        while slug in slug_set:
            stats['duplicates_handled'] += 1
            slug = f"{base_slug}-{idx}"
            idx += 1
        slug_set.add(slug)

        services.append({
            'id': slug,
            'slug': slug,
            'service_code': f"GEN-{idx}",
            'name': name,
            'category_id': 'genetics',
            'category_name': 'Genomics & Genetics',
            'subcategory': cat_raw,
            'provider': 'Health Express Genetics',
            'description': desc,
            'overview': desc,
            'mrp': mrp,
            'price': mrp,
            'selling_price': mrp,
            'discount_price': mrp,
            'b2b_price': None,
            'discount_percentage': None,
            'turnaround_time': tat or '7-14 Days',
            'patient_preparation': prep or 'Saliva/Blood collection kit provided',
            'preparation': prep or 'Saliva/Blood collection kit provided',
            'specimen_type': sample or 'Saliva / Blood',
            'sample_type': sample or 'Saliva / Blood',
            'home_collection_available': True,
            'centre_visit_required': False,
            'service_type': 'genomics',
            'source': 'Genomics'
        })
        stats['genomics'] += 1

# 4. IMPORT HEALTH PACKAGES (PREVENTIVE CARE + SEXUAL WELLNESS MATRIX)
f4 = downloads_dir / 'HealthPackages.xlsx'
if f4.exists():
    wb4 = openpyxl.load_workbook(f4, data_only=True)
    
    # Sheet 1: Preventive Care
    if 'Preventive Care' in wb4.sheetnames:
        sp = wb4['Preventive Care']
        rows_sp = list(sp.iter_rows(values_only=True))
        header_p = rows_sp[2]
        mrp_row_p = rows_sp[43]

        for col_idx in range(3, 12):
            pkg_name = str(header_p[col_idx]).strip() if header_p[col_idx] else None
            if not pkg_name: continue
            
            mrp_val = mrp_row_p[col_idx] if col_idx < len(mrp_row_p) else None
            mrp = float(mrp_val) if mrp_val and str(mrp_val).replace('.', '', 1).isdigit() else 2999
            
            params = []
            for r_idx in range(3, 43):
                r = rows_sp[r_idx]
                test_name = r[1]
                val = str(r[col_idx]).strip().lower() if col_idx < len(r) and r[col_idx] is not None else ''
                if test_name and val == 'yes':
                    params.append(str(test_name).strip())
                    
            base_slug = slugify(f"{pkg_name}-preventive")
            slug = base_slug
            idx = 1
            while slug in slug_set:
                slug = f"{base_slug}-{idx}"
                idx += 1
            slug_set.add(slug)

            services.append({
                'id': slug,
                'slug': slug,
                'service_code': f"PKG-PREV-{col_idx}",
                'name': pkg_name,
                'category_id': 'health-packages',
                'category_name': 'Health Packages',
                'subcategory': 'Preventive Health',
                'provider': 'Health Express Network',
                'description': f"Comprehensive preventive checkup panel containing {len(params)} diagnostic parameters.",
                'overview': f"Comprehensive preventive checkup panel containing {len(params)} diagnostic parameters.",
                'mrp': mrp,
                'price': mrp,
                'selling_price': mrp,
                'discount_price': mrp,
                'b2b_price': None,
                'discount_percentage': None,
                'turnaround_time': '24 Hours',
                'patient_preparation': '10-12 hours fasting required before sample collection.',
                'preparation': '10-12 hours fasting required before sample collection.',
                'specimen_type': 'Blood & Urine',
                'sample_type': 'Blood & Urine',
                'home_collection_available': True,
                'centre_visit_required': False,
                'parameters': params,
                'parameters_count': len(params),
                'service_type': 'package',
                'source': 'Preventive Care'
            })
            stats['packages'] += 1

    # Sheet 2: SexualWellness
    if 'SexualWellness' in wb4.sheetnames:
        sw = wb4['SexualWellness']
        rows_sw = list(sw.iter_rows(values_only=True))
        header_group = rows_sw[0]
        header_name = rows_sw[1]
        mrp_row_sw = rows_sw[38]

        for col_idx in range(2, 9):
            group_title = ''
            for prev in range(col_idx, -1, -1):
                if header_group[prev]:
                    group_title = str(header_group[prev]).strip()
                    break
                    
            pkg_name = str(header_name[col_idx]).strip() if header_name[col_idx] else f"Package {col_idx}"
            mrp_val = mrp_row_sw[col_idx] if col_idx < len(mrp_row_sw) else None
            mrp = float(mrp_val) if mrp_val and str(mrp_val).replace('.', '', 1).isdigit() else 4999
            
            params = []
            for r_idx in range(2, 38):
                r = rows_sw[r_idx]
                test_name = r[1]
                val = str(r[col_idx]).strip().lower() if col_idx < len(r) and r[col_idx] is not None else ''
                if test_name and val == 'yes':
                    params.append(str(test_name).strip())
                    
            clean_group = group_title.replace('Packages', '').replace('(Premium)', '').strip()
            full_name = f"{clean_group} - {pkg_name}"
            if pkg_name.lower() in clean_group.lower():
                full_name = clean_group

            base_slug = slugify(f"{full_name}-sexual-wellness")
            slug = base_slug
            idx = 1
            while slug in slug_set:
                slug = f"{base_slug}-{idx}"
                idx += 1
            slug_set.add(slug)

            services.append({
                'id': slug,
                'slug': slug,
                'service_code': f"PKG-SEX-{col_idx}",
                'name': full_name,
                'category_id': 'health-packages',
                'category_name': 'Health Packages',
                'subcategory': 'Sexual Wellness',
                'provider': 'Health Express Network',
                'description': f"Specialized sexual health and fertility package including {len(params)} diagnostic parameters & specialist consultation.",
                'overview': f"Specialized sexual health and fertility package including {len(params)} diagnostic parameters & specialist consultation.",
                'mrp': mrp,
                'price': mrp,
                'selling_price': mrp,
                'discount_price': mrp,
                'b2b_price': None,
                'discount_percentage': None,
                'turnaround_time': '24-48 Hours',
                'patient_preparation': 'Confidential home sample pickup. No fasting required unless specified.',
                'preparation': 'Confidential home sample pickup. No fasting required unless specified.',
                'specimen_type': 'Blood & Urine',
                'sample_type': 'Blood & Urine',
                'home_collection_available': True,
                'centre_visit_required': False,
                'parameters': params,
                'parameters_count': len(params),
                'service_type': 'package',
                'source': 'Sexual Wellness'
            })
            stats['packages'] += 1

print("=== PARSING & NORMALIZATION REPORT ===")
print(f"Total Normalized Services: {len(services)}")
print(f"LabSpring Records: {stats['labspring']}")
print(f"Imaging Records: {stats['imaging']}")
print(f"Genomics Records: {stats['genomics']}")
print(f"Health Package Records: {stats['packages']}")

# Write to src/data/services.js
categories_map = {
    'lab-tests': {'id': 'lab-tests', 'name': 'Lab Tests', 'icon': 'FlaskConical', 'count': 0},
    'imaging': {'id': 'imaging', 'name': 'Imaging & Radiology', 'icon': 'Scan', 'count': 0},
    'genetics': {'id': 'genetics', 'name': 'Genomics & Genetics', 'icon': 'Dna', 'count': 0},
    'health-packages': {'id': 'health-packages', 'name': 'Health Packages', 'icon': 'ShieldCheck', 'count': 0},
}

providers_set = set()
for s in services:
    cat_id = s['category_id']
    if cat_id in categories_map:
        categories_map[cat_id]['count'] += 1
    if s.get('provider'):
        providers_set.add(s['provider'])

categories_list = list(categories_map.values())
providers_list = sorted(list(providers_set))

# Output JS data file
js_content = f"""// Normalized Health Express Service Catalog ({len(services)} Services)
// Auto-generated from client Excel datasets: LabSpring (1764), Imaging (417), Genomics (10), HealthPackages (16)

export const CATEGORIES = {json.dumps(categories_list, indent=2)};

export const PROVIDERS = {json.dumps(providers_list, indent=2)};

export const ALL_SERVICES = {json.dumps(services, indent=2)};
"""

target_js = Path(r'c:\desktop\healthexpress-website\src\data\services.js')
target_js.write_text(js_content, encoding='utf-8')
print(f"\nWritten {len(services)} services to {target_js}")

# Output SQL migration file
sql_lines = [
    "-- ============================================================",
    "-- HEALTH EXPRESS — REAL SERVICE CATALOG PRODUCTION MIGRATION",
    f"-- Total Normalized Services: {len(services)} records",
    "-- ============================================================",
    "",
    "CREATE TABLE IF NOT EXISTS public.services (",
    "  id UUID DEFAULT gen_random_uuid() PRIMARY KEY,",
    "  slug TEXT NOT NULL UNIQUE,",
    "  service_code TEXT NULL,",
    "  service_name TEXT NOT NULL,",
    "  category_id TEXT NOT NULL,",
    "  category_name TEXT NOT NULL,",
    "  subcategory TEXT NULL,",
    "  provider TEXT NULL,",
    "  description TEXT NULL,",
    "  overview TEXT NULL,",
    "  mrp NUMERIC(10,2) NULL,",
    "  selling_price NUMERIC(10,2) NOT NULL,",
    "  b2b_price NUMERIC(10,2) NULL,",
    "  discount_percentage TEXT NULL,",
    "  turnaround_time TEXT NULL,",
    "  patient_preparation TEXT NULL,",
    "  specimen_type TEXT NULL,",
    "  home_collection_available BOOLEAN DEFAULT false,",
    "  centre_visit_required BOOLEAN DEFAULT false,",
    "  parameters JSONB DEFAULT '[]'::jsonb,",
    "  parameters_count INTEGER DEFAULT 0,",
    "  service_type TEXT DEFAULT 'lab',",
    "  source TEXT DEFAULT 'catalog',",
    "  active BOOLEAN DEFAULT true,",
    "  created_at TIMESTAMPTZ DEFAULT now() NOT NULL,",
    "  updated_at TIMESTAMPTZ DEFAULT now() NOT NULL",
    ");",
    "",
    "CREATE INDEX IF NOT EXISTS idx_services_slug ON public.services(slug);",
    "CREATE INDEX IF NOT EXISTS idx_services_code ON public.services(service_code);",
    "CREATE INDEX IF NOT EXISTS idx_services_category ON public.services(category_id);",
    "CREATE INDEX IF NOT EXISTS idx_services_provider ON public.services(provider);",
    "CREATE INDEX IF NOT EXISTS idx_services_price ON public.services(selling_price);",
    "CREATE INDEX IF NOT EXISTS idx_services_active ON public.services(active);",
    "",
    "ALTER TABLE public.services ENABLE ROW LEVEL SECURITY;",
    "",
    "DROP POLICY IF EXISTS \"Public can view active services\" ON public.services;",
    "CREATE POLICY \"Public can view active services\" ON public.services FOR SELECT USING (active = true);",
    "",
    "DROP POLICY IF EXISTS \"Admin can manage services\" ON public.services;",
    "CREATE POLICY \"Admin can manage services\" ON public.services FOR ALL USING (public.check_is_admin() = true);",
    ""
]

def sql_quote(val):
    if val is None: return "NULL"
    if isinstance(val, bool): return "TRUE" if val else "FALSE"
    if isinstance(val, (int, float)): return str(val)
    if isinstance(val, (list, dict)): return f"'{json.dumps(val)}'::jsonb"
    escaped = str(val).replace("'", "''")
    return f"'{escaped}'"

for s in services:
    mrp_q = sql_quote(s.get('mrp'))
    selling_q = sql_quote(s.get('selling_price'))
    b2b_q = sql_quote(s.get('b2b_price'))
    code_q = sql_quote(s.get('service_code'))
    desc_q = sql_quote(s.get('description'))
    overview_q = sql_quote(s.get('overview'))
    sub_q = sql_quote(s.get('subcategory'))
    tat_q = sql_quote(s.get('turnaround_time'))
    prep_q = sql_quote(s.get('patient_preparation'))
    sample_q = sql_quote(s.get('specimen_type'))
    disc_pct_q = sql_quote(s.get('discount_percentage'))
    hc_q = sql_quote(s.get('home_collection_available'))
    cv_q = sql_quote(s.get('centre_visit_required'))
    params_q = sql_quote(s.get('parameters', []))
    params_cnt = s.get('parameters_count', 0)

    sql_lines.append(f"""INSERT INTO public.services (
  slug, service_code, service_name, category_id, category_name, subcategory, provider,
  description, overview, mrp, selling_price, b2b_price, discount_percentage,
  turnaround_time, patient_preparation, specimen_type, home_collection_available,
  centre_visit_required, parameters, parameters_count, service_type, source
) VALUES (
  {sql_quote(s['slug'])}, {code_q}, {sql_quote(s['name'])}, {sql_quote(s['category_id'])}, {sql_quote(s['category_name'])}, {sub_q}, {sql_quote(s['provider'])},
  {desc_q}, {overview_q}, {mrp_q}, {selling_q}, {b2b_q}, {disc_pct_q},
  {tat_q}, {prep_q}, {sample_q}, {hc_q},
  {cv_q}, {params_q}, {params_cnt}, {sql_quote(s['service_type'])}, {sql_quote(s['source'])}
) ON CONFLICT (slug) DO UPDATE SET
  service_name = EXCLUDED.service_name,
  selling_price = EXCLUDED.selling_price,
  mrp = EXCLUDED.mrp,
  turnaround_time = EXCLUDED.turnaround_time,
  updated_at = now();""")

target_sql = Path(r'c:\desktop\healthexpress-website\supabase\migrations\20261001_services_catalog.sql')
target_sql.write_text("\n".join(sql_lines), encoding='utf-8')
print(f"Written SQL migration to {target_sql}")
